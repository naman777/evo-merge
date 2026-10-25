"""Build the project report (.docx) from report.md.

    python report/build_docx.py            -> report/Soft_Computing_Project_Report.docx

Handles the Markdown subset report.md uses: #/##/### headings, paragraphs with **bold**, *italic*,
`code` and [links](url), pipe tables, numbered and bullet lists, ![caption](image) figures,
$$...$$ display formulas (rendered with matplotlib) and fenced code blocks.
"""
import os
import re
import tempfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402,F401
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Inches, Pt, RGBColor  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "report.md")
OUT = os.path.join(HERE, "Soft_Computing_Project_Report.docx")
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|[θτ]_[a-z]+)")
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FAF")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_hyperlink(paragraph, label, url)
            continue
        elif re.fullmatch(r"[θτ]_[a-z]+", part):
            paragraph.add_run(part[0])
            run = paragraph.add_run(part[2:])
            run.font.subscript = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = size


def shade(cell, hex_fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, cells in enumerate([header] + body):
        for c, text in enumerate(cells):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text.replace("\\[", "[").replace("\\]", "]"), size=Pt(9.5))
            if r == 0:
                for run in p.runs:
                    run.bold = True
                shade(cell, "DCE6F1")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_formula(tex, tmpdir, n):
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${tex}$", fontsize=12)
    path = os.path.join(tmpdir, f"formula{n}.png")
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.05, transparent=False, facecolor="white")
    plt.close(fig)
    return path


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def setup(doc):
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [("Heading 1", 15), ("Heading 2", 12.5)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)
    add_page_number(section)


def title_page(doc, title, lines):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = ACCENT
    doc.add_paragraph()
    return lines


def build():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()
    doc = Document()
    setup(doc)
    tmpdir = tempfile.mkdtemp()
    n_formula = 0
    i = 0
    on_title_page = True
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            title_page(doc, line[2:], None)
        elif line.startswith("## "):
            if on_title_page:  # everything before the first section heading is the title page
                p = doc.add_paragraph()
                p.add_run().add_break(WD_BREAK.PAGE)
                on_title_page = False
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= {"-", ":", " "} for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith("$$"):
            n_formula += 1
            path = render_formula(line.strip("$ "), tmpdir, n_formula)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path)  # natural size: 300 dpi at 12 pt
        elif line.startswith("!["):
            caption, path = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line).groups()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.normpath(os.path.join(HERE, path)), width=Inches(5.8))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(caption)
            run.italic = True
            run.font.size = Pt(9.5)
        elif line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.5)
                run = p.add_run(lines[i])
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                i += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\. ", line) or line.startswith("- "):
            numbered = bool(re.match(r"^\d+\. ", line))
            marker, text = (re.match(r"^(\d+\.) (.*)$", line).groups() if numbered else ("•", line[2:]))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            p.add_run(marker + "\t")
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            add_inline(p, text)
        else:
            p = doc.add_paragraph()
            if on_title_page:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line.replace("\\[", "[").replace("\\]", "]"))
        i += 1
    doc.core_properties.title = "Evolutionary Optimization of Model Merging Recipes for Hindi Math Reasoning"
    doc.core_properties.author = "Naman Kundra"
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
